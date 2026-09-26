from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
import zipfile
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import toa_web
from toa_web import (
    RequestError,
    ToAProtocolRequest,
    ToAWebHandler,
    _job_public,
    _new_job,
    _output_files,
    _public_status,
    _read_json,
    _safe_child,
    _start_process,
    _update_job,
    _validate_docx,
    _validate_pdf,
    _validate_settings,
    create_server,
)


class ToAWebTests(unittest.TestCase):
    def test_plugin_transport_reuses_the_web_handlers_without_a_server(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary) / "state"
            inbox = Path(temporary) / "inbox"
            inbox.mkdir()
            request = ToAProtocolRequest(root, inbox, "/api/jobs")
            ToAWebHandler.do_POST(request)
            self.assertEqual(request.response["status"], 201)
            job = request.response["body"]
            self.assertTrue((root / "jobs" / job["id"] / "job.json").is_file())

    def test_build_processes_have_a_global_concurrency_limit(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            _, job = _new_job(Path(temporary), "brief.docx")
            with mock.patch.object(toa_web, "_RUNNING", {"first", "second"}):
                with self.assertRaises(RequestError) as error:
                    _start_process(job, "build", [sys.executable, "--version"])
            self.assertEqual(error.exception.status, 503)
            self.assertEqual(_read_json(job / "job.json", {})["state"], "ready")

    def test_request_log_omits_job_ids_queries_and_filenames(self) -> None:
        handler = object.__new__(ToAWebHandler)
        handler.client_address = ("127.0.0.1", 12345)
        handler.command = "GET"
        handler.path = "/api/jobs/job-secret/files/Client-Matter.pdf?token=secret"
        with mock.patch("sys.stderr", new_callable=io.StringIO) as output:
            handler.log_message('"%s" %s %s', handler.path, "200", "-")
        logged = output.getvalue()
        self.assertIn("GET /api/jobs/* 200", logged)
        self.assertNotIn("job-secret", logged)
        self.assertNotIn("Client-Matter", logged)
        self.assertNotIn("token", logged)

    def test_upload_validators_and_traversal(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            docx = root / "valid.docx"
            with zipfile.ZipFile(docx, "w") as archive:
                archive.writestr("word/document.xml", "<document/>")
            pdf = root / "valid.pdf"
            pdf.write_bytes(b"%PDF-1.7\n%%EOF")
            _validate_docx(docx)
            _validate_pdf(pdf)
            with zipfile.ZipFile(docx, "a") as archive:
                archive.writestr("word/vbaProject.bin", b"untrusted")
            with self.assertRaises(RequestError):
                _validate_docx(docx)
            with self.assertRaises(RequestError):
                _safe_child(root, "../outside")
            with self.assertRaises(RequestError):
                _safe_child(root, r"outputs\..\outside")

    def test_status_and_blank_manual_job(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            server = create_server(0, state_root=Path(temporary))
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base = f"http://127.0.0.1:{server.server_port}"
            try:
                with urllib.request.urlopen(f"{base}/api/status") as response:
                    self.assertEqual(response.headers["Server"], "AuthoritiesHelper")
                    status = json.load(response)
                self.assertEqual(status["service"], "authorities-helper")
                with mock.patch(
                    "toa_web._dependency_status",
                    side_effect=RuntimeError("secret local path"),
                ):
                    with self.assertRaises(urllib.error.HTTPError) as error:
                        urllib.request.urlopen(f"{base}/api/status")
                    self.assertEqual(error.exception.code, 500)
                    self.assertEqual(
                        json.load(error.exception),
                        {"error": "Internal server error."},
                    )
                with urllib.request.urlopen(f"{base}/styles.css") as response:
                    styles = response.read().decode("utf-8")
                self.assertIn("--brand: #d52b1e;", styles)
                self.assertIn("--brand-dark: #b72016;", styles)
                self.assertIn("--paper: #f3f4f6;", styles)
                self.assertIn("--control: #111827;", styles)
                self.assertIn("@media (max-width: 480px)", styles)
                self.assertIn("grid-template-rows: 52px 45px minmax(0, 1fr);", styles)
                self.assertIn(".workflow-primary", styles)
                self.assertIn(".file-button:focus-within", styles)
                self.assertIn("overflow: hidden;", styles)
                self.assertIn("#automatic.view.active", styles)
                self.assertNotIn("overflow-wrap: anywhere", styles)
                self.assertNotIn("filter: brightness", styles)
                with urllib.request.urlopen(f"{base}/index.html") as response:
                    csp = response.headers["Content-Security-Policy"]
                    self.assertEqual(response.headers["Referrer-Policy"], "no-referrer")
                    self.assertIn("camera=()", response.headers["Permissions-Policy"])
                    index = response.read().decode("utf-8")
                self.assertIn("base-uri 'none'", csp)
                self.assertIn("object-src 'none'", csp)
                with urllib.request.urlopen(f"{base}/app.js") as response:
                    app = response.read().decode("utf-8")
                with urllib.request.urlopen(f"{base}/?mode=mike") as response:
                    embedded = response.read().decode("utf-8")
                self.assertIn('<html lang="en" class="mike-mode">', embedded)
                self.assertIn('classList.toggle("mike-mode", mikeMode)', app)
                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(f"{base}/mode.js")
                self.assertEqual(error.exception.code, 404)
                self.assertIn(
                    'data-view="automatic" aria-current="page">Automatic',
                    index,
                )
                self.assertIn('data-view="manual">Manual', index)
                self.assertNotIn('id="manual-open"', index)
                self.assertNotIn('id="manual-back"', index)
                self.assertNotIn("Start with PDFs", index)
                self.assertNotIn("Use a Word document", index)
                self.assertEqual(index.count('id="job-status"'), 1)
                self.assertEqual(index.count('id="job-progress"'), 1)
                self.assertNotIn("status-strip", index)
                self.assertIn('data-view="history"', index)
                self.assertIn('id="history-list"', index)
                self.assertNotIn('id="history-select"', index)
                self.assertNotIn("data-step=", index)
                self.assertIn(
                    'id="document-input" type="file" accept=".docx,.pdf,',
                    index,
                )
                self.assertLess(index.index('id="document-input"'), index.index('id="review-card"'))
                self.assertLess(index.index('id="review-card"'), index.index('id="build-card"'))
                self.assertLess(index.index('id="build-card"'), index.index('id="insert-card"'))
                self.assertLess(index.index('id="insert-card"'), index.index('id="output-card"'))
                self.assertIn('id="active-document-name"', index)
                self.assertIn('id="replace-document"', index)
                self.assertNotIn('id="review-save"', index)
                self.assertNotIn("Needs review", index)
                self.assertNotIn("Reviewed</", index)
                self.assertIn('id="scanned-pdf-dialog"', index)
                self.assertIn('id="setup-dialog"', index)
                self.assertIn('id="setup-open"', index)
                self.assertIn("Change workflow defaults", index)
                self.assertIn('id="setup-submit"', index)
                self.assertNotIn('id="settings-save"', index)
                self.assertIn("Use originals; rebuild missing sources", index)
                self.assertIn("Use originals; add a page for missing sources", index)
                self.assertIn("Rebuild all sources from text", index)
                self.assertIn('id="setup-marking-fieldset"', index)
                self.assertEqual(index.count('name="setup-highlight-style"'), 5)
                self.assertIn("Black paragraph line", index)
                self.assertIn("Right-margin marker + exact quote", index)
                self.assertIn("Highlight the whole cited paragraph", index)
                self.assertIn("Highlight exact quotes only", index)
                self.assertIn("No passage marks", index)
                self.assertNotIn("Local runtime", index)
                self.assertNotIn('id="runtime-status"', index)
                self.assertNotIn('id="mike-session"', index)
                self.assertNotIn('id="job-operation"', index)
                self.assertNotIn('id="job-percent"', index)
                self.assertNotIn('id="job-error"', index)
                self.assertNotIn("No document imported", index)
                self.assertNotIn("Maximum 64 MB", index)
                self.assertNotIn(
                    "These choices are used by the automatic workflow.",
                    index,
                )
                self.assertNotIn("localStorage", app)
                self.assertIn("const SETUP_VERSION = 3;", app)
                self.assertIn("settings.setup_version >= SETUP_VERSION", app)
                self.assertNotIn('query.get("setup")', app)
                self.assertNotIn('name="setup-workflow"', index)
                self.assertIn('openSetup("automatic")', app)
                self.assertIn('openSetup("manual")', app)
                self.assertIn('start ? "Start" : "Done"', app)
                self.assertNotIn("Needs attention", app)
                self.assertNotIn("historyState", app)
                self.assertNotIn('$("#settings-save")', app)
                self.assertGreaterEqual(app.count("saveSettings(false);"), 3)
                self.assertIn(
                    "settings.highlight_style = dialog.querySelector"
                    "('input[name=\"setup-highlight-style\"]:checked').value;",
                    app,
                )
                self.assertIn("historyJobsPromise = api(jobsPath())", app)
                self.assertIn('view === "history"', app)
                self.assertIn("async function openHistoryJob(job)", app)
                boot = app[app.index("async function boot()") :]
                self.assertIn("void preloadHistory()", boot)
                self.assertIn("settingsReady = loadSettings();", boot)
                self.assertIn("if (currentJob) {\n    await settingsReady;", boot)
                self.assertNotIn("await loadSettings();", boot)
                self.assertIn('"mike:authorities-helper-ready"', boot)
                self.assertIn('"mike:authorities-helper-error"', boot)
                self.assertIn("attempt: readyAttempt", boot)
                self.assertNotIn('api("/api/status")', app)
                self.assertNotIn("data_root", app)
                self.assertIn("sessionStorage.getItem(sessionKey)", app)
                self.assertIn('sessionStorage.getItem("toa-session")', app)
                self.assertIn("requestedJob || (!mikeMode", app)
                self.assertIn('"mike:authorities-helper-probe"', app)
                self.assertIn("if (!bootComplete || !parentOrigin) return;", app)
                self.assertIn("const requestedJob = mikeMode &&", app)
                self.assertIn("if (job.has_review) await loadReview();", app)
                self.assertIn('"Local only"', app)
                self.assertIn('"Resolve names"', app)
                self.assertIn("Scanned source pages", index)
                self.assertIn('if (!accepted) return false;', app)
                self.assertIn('settings.prompt_for_scanned_pdfs = !$("#scanned-do-not-show").checked;', app)
                self.assertIn('output_mode: "book"', app)
                self.assertIn('"Book of Authorities (PDF)"', app)
                self.assertIn('"Table of Authorities (Word)"', app)
                self.assertNotIn('"Build table"', app)
                self.assertNotIn('"Build both"', app)
                self.assertIn("enrich_spans: true", app)
                self.assertIn('["output_mode", "Create"', app)
                self.assertNotIn("setting-info", app)
                self.assertNotIn("setting-tooltip", app)
                self.assertIn("syncBuildFields();", app)
                self.assertIn('button.setAttribute("aria-current", "page")', app)
                self.assertIn('className = "citation-item"', app)
                self.assertIn("function renderReviewEditor()", app)
                self.assertIn('id="review-surface"', app)
                self.assertIn("Use selection as authority", app)
                self.assertIn("Use selection as pinpoint", app)
                self.assertIn("Split at cursor", app)
                self.assertIn("Merge with previous", app)
                self.assertIn("/review/action", app)
                self.assertNotIn('data-field="kind"', app)
                self.assertNotIn('data-field="citation"', app)
                self.assertIn("Open CanLII PDF", app)
                self.assertIn("Add downloaded PDF", app)
                self.assertIn("CanLII opens separately", app)
                with urllib.request.urlopen(f"{base}/court-profiles.json") as response:
                    profiles = json.load(response)
                self.assertEqual(
                    {"general", "abkb", "abca", "fc", "fca"},
                    {item["id"] for item in profiles["profiles"]},
                )
                self.assertIn('class="scan-flow"', index)
                self.assertNotIn("<table", index)
                self.assertNotIn("ReviewState JSON", app)
                self.assertNotIn("switchStep", app)

                request = urllib.request.Request(
                    f"{base}/api/jobs",
                    data=b"",
                    method="POST",
                )
                with urllib.request.urlopen(request) as response:
                    job = json.load(response)
                self.assertRegex(job["id"], r"^[0-9a-f]{32}$")
                same_origin = urllib.request.Request(
                    f"{base}/api/jobs",
                    data=b"",
                    method="POST",
                    headers={"Origin": base},
                )
                with urllib.request.urlopen(same_origin) as response:
                    self.assertEqual(response.status, 201)
                foreign_port = urllib.request.Request(
                    f"{base}/api/jobs",
                    data=b"",
                    method="POST",
                    headers={"Origin": "http://127.0.0.1:1"},
                )
                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(foreign_port)
                self.assertEqual(error.exception.code, 403)
                manual = json.loads(
                    (Path(temporary) / "jobs" / job["id"] / "manual" / "state.json").read_text(
                        encoding="utf-8"
                    )
                )
                self.assertEqual(manual["book_title"], "Book of Authorities")
                internal = Path(temporary) / "jobs" / job["id"] / "outputs" / "build" / "authorities"
                internal.mkdir(parents=True)
                (internal / "tab-1.pdf").write_bytes(b"%PDF-1.7\n%%EOF")
                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(
                        f"{base}/api/jobs/{job['id']}/files/outputs/build/authorities/tab-1.pdf"
                    )
                self.assertEqual(error.exception.code, 403)

                bad = urllib.request.Request(
                    f"{base}/api/status",
                    headers={"Host": "example.com"},
                )
                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(bad)
                self.assertEqual(error.exception.code, 403)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=2)

    def test_jobs_can_be_scoped_to_a_project(self) -> None:
        first = "11111111-1111-4111-8111-111111111111"
        second = "22222222-2222-4222-8222-222222222222"
        with tempfile.TemporaryDirectory() as temporary:
            server = create_server(0, state_root=Path(temporary))
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base = f"http://127.0.0.1:{server.server_port}"
            try:
                for project in (first, second):
                    request = urllib.request.Request(
                        f"{base}/api/jobs?project={project}",
                        data=b"",
                        method="POST",
                    )
                    with urllib.request.urlopen(request) as response:
                        job = json.load(response)
                    self.assertEqual(job["project_id"], project)

                with urllib.request.urlopen(f"{base}/api/jobs?project={first}") as response:
                    scoped = json.load(response)["jobs"]
                self.assertEqual(len(scoped), 1)
                self.assertEqual(scoped[0]["project_id"], first)

                with urllib.request.urlopen(f"{base}/api/jobs") as response:
                    all_jobs = json.load(response)["jobs"]
                self.assertEqual(len(all_jobs), 2)

                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(f"{base}/api/jobs?project=not-a-project")
                self.assertEqual(error.exception.code, 400)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=2)

    def test_public_files_are_manifest_deliverables_not_build_internals(self) -> None:
        expected = {
            "book": {"brief.book-of-authorities.pdf"},
            "table": {"brief.updated.docx"},
            "both": {"brief.book-of-authorities.pdf", "brief.updated.docx"},
        }
        for mode, names in expected.items():
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as temporary:
                job = Path(temporary) / ("a" * 32)
                build = job / "outputs" / "build"
                internal = build / "authorities"
                manual = job / "manual" / "outputs"
                internal.mkdir(parents=True)
                manual.mkdir(parents=True)
                assembled = build / "brief.book-of-authorities.pdf"
                table = build / "brief.updated.docx"
                source_pdf = build / "brief.document.pdf"
                tab_pdf = internal / "tab-1.pdf"
                manual_book = manual / "old-manual-book.pdf"
                latest_manual_book = manual / "current-manual-book.pdf"
                for path in (
                    assembled,
                    table,
                    source_pdf,
                    tab_pdf,
                    manual_book,
                    latest_manual_book,
                ):
                    path.write_bytes(b"x")
                (job / "job.json").write_text(
                    json.dumps({"operation": "finalization", "state": "complete"}),
                    encoding="utf-8",
                )
                (build / "brief.toa-manifest.json").write_text(
                    json.dumps(
                        {
                            "output_mode": mode,
                            "outputs": {
                                "book_of_authorities_pdf": str(assembled),
                                "annotated_docx": str(table),
                                "source_document_pdf": str(source_pdf),
                            },
                        }
                    ),
                    encoding="utf-8",
                )
                self.assertEqual({row["name"] for row in _output_files(job)}, names)
                public = _job_public(job)
                self.assertEqual(public["output_mode"], mode)
                self.assertEqual(
                    public["message"],
                    {
                        "book": "Book ready",
                        "table": "Table ready",
                        "both": "Book and table ready",
                    }[mode],
                )
                self.assertEqual({row["name"] for row in public["files"]}, names)

                (job / "job.json").write_text(
                    json.dumps(
                        {
                            "operation": "manual book",
                            "manual_output": str(latest_manual_book),
                        }
                    ),
                    encoding="utf-8",
                )
                self.assertEqual(
                    {row["name"] for row in _output_files(job)},
                    {"current-manual-book.pdf"},
                )

    def test_review_edits_are_small_validated_and_saved_automatically(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            job_id, job = _new_job(root, "factum.docx")
            review = {
                "parts": [
                    {
                        "uid": "part-1",
                        "unit_key": "footnote:1",
                        "index": 1,
                        "kind": "case",
                        "citation": "R v Grant, 2009 SCC 32",
                        "pinpoint_fragments": [],
                    },
                    {
                        "uid": "part-2",
                        "unit_key": "footnote:2",
                        "index": 1,
                        "kind": "case",
                        "citation": "R v Jordan, 2016 SCC 27",
                        "pinpoint_fragments": [],
                    },
                ]
            }
            (job / "review.json").write_text(json.dumps(review), encoding="utf-8")
            server = create_server(0, state_root=root)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base = f"http://127.0.0.1:{server.server_port}"
            edit = {
                "edits": [
                    {
                        "part_id": "part-1",
                        "changes": {
                            "kind": "statute",
                            "citation": "Criminal Code, RSC 1985, c C-46",
                            "pinpoint_fragments": ["s 7"],
                        },
                    }
                ]
            }
            try:
                request = urllib.request.Request(
                    f"{base}/api/jobs/{job_id}/review",
                    data=json.dumps(edit).encode(),
                    headers={"Content-Type": "application/json"},
                    method="PUT",
                )
                with urllib.request.urlopen(request) as response:
                    self.assertEqual(json.load(response), {"saved": 1})
                saved = _read_json(job / "review.json", {})
                self.assertEqual(saved["parts"][0]["kind"], "statute")
                self.assertEqual(saved["parts"][0]["pinpoint_fragments"], ["s 7"])
                self.assertTrue(saved["parts"][0]["reviewed"])
                self.assertEqual(saved["parts"][1], review["parts"][1])
                self.assertLess(len(json.dumps(edit)), len(json.dumps(review)))
                if os.environ.get("TOA_PRINT_METRICS"):
                    print(
                        "review-save-bytes",
                        json.dumps(
                            {
                                "before": len(json.dumps(review)),
                                "after": len(json.dumps(edit)),
                            }
                        ),
                    )

                legacy = urllib.request.Request(
                    f"{base}/api/jobs/{job_id}/review",
                    data=json.dumps(review).encode(),
                    headers={"Content-Type": "application/json"},
                    method="PUT",
                )
                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(legacy)
                self.assertEqual(error.exception.code, 400)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=2)

    def test_review_surface_actions_match_the_python_workflow(self) -> None:
        from toa_maker import (
            DeterministicPart,
            ReviewState,
            TextUnit,
            _anchor_spans,
            _part_from,
        )

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            job_id, job = _new_job(root, "factum.docx")
            text = (
                "R v Grant, 2009 SCC 32 at para 25; "
                "R v Jordan, 2016 SCC 27 at para 47."
            )
            unit = TextUnit("footnote:8", "footnote", 1, 8, text)
            part = _part_from(
                unit,
                1,
                "manual",
                [],
                DeterministicPart(
                    0,
                    len(text),
                    text,
                    tuple(item[2] for item in _anchor_spans(text)),
                ),
            )
            part.uid = "footnote:8:p1"
            ReviewState("factum.docx", [unit], [part], "now").save(
                job / "review.json",
                compact=True,
            )
            server = create_server(0, state_root=root)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base = f"http://127.0.0.1:{server.server_port}"

            def action(payload: dict[str, object]) -> dict[str, object]:
                request = urllib.request.Request(
                    f"{base}/api/jobs/{job_id}/review/action",
                    data=json.dumps(payload).encode(),
                    headers={"Content-Type": "application/json"},
                    method="POST",
                )
                with urllib.request.urlopen(request) as response:
                    return json.load(response)

            try:
                cut = text.index("R v Jordan")
                split = action(
                    {
                        "action": "split",
                        "part_id": part.part_id,
                        "cursor": cut,
                    }
                )
                self.assertEqual(len(split["review"]["parts"]), 2)
                self.assertEqual(
                    [item["text"] for item in split["review"]["parts"]],
                    [
                        "R v Grant, 2009 SCC 32 at para 25;",
                        "R v Jordan, 2016 SCC 27 at para 47.",
                    ],
                )

                merged = action(
                    {
                        "action": "merge",
                        "part_id": split["review"]["parts"][1]["uid"],
                    }
                )
                self.assertEqual(len(merged["review"]["parts"]), 1)
                merged_id = merged["selected_part_id"]

                grant_end = text.index(" at para 25")
                action(
                    {
                        "action": "authority",
                        "part_id": merged_id,
                        "start": 0,
                        "end": grant_end,
                    }
                )
                pinpoint_start = text.index("para 25")
                marked = action(
                    {
                        "action": "pinpoint",
                        "part_id": merged_id,
                        "start": pinpoint_start,
                        "end": pinpoint_start + len("para 25"),
                    }
                )
                saved = marked["review"]["parts"][0]
                self.assertEqual(saved["authority_start"], 0)
                self.assertEqual(saved["authority_end"], grant_end)
                self.assertEqual(saved["pinpoint_fragments"], ["par25"])
                self.assertTrue(saved["reviewed"])

                with self.assertRaises(urllib.error.HTTPError) as error:
                    action(
                        {
                            "action": "split",
                            "part_id": merged_id,
                            "cursor": 0,
                        }
                    )
                self.assertEqual(error.exception.code, 400)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=2)

    def test_progress_is_utf8_throttled_and_public(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            _, job = _new_job(root, "mémoire.docx")
            command = [
                sys.executable,
                "-X",
                "utf8",
                "-c",
                "[print(f'PROGRESS\\t{i}\\tRésumé — PDFs') for i in range(100)]",
            ]
            with mock.patch("toa_web._update_job", wraps=_update_job) as update:
                _start_process(job, "build", command)
                for _ in range(100):
                    state = _read_json(job / "job.json", {})
                    if state.get("state") != "running":
                        break
                    time.sleep(0.02)
            self.assertEqual(state["state"], "complete", state)
            self.assertEqual(state["log"][-1], "PROGRESS\t99\tRésumé — PDFs")
            self.assertLessEqual(update.call_count, 5)
            if os.environ.get("TOA_PRINT_METRICS"):
                print(
                    "progress-writes",
                    json.dumps({"before": 102, "after": update.call_count}),
                )
            public = _job_public(job)
            self.assertEqual(public["message"], "Book ready")
            self.assertNotIn("log", public)
            if os.environ.get("TOA_PRINT_METRICS"):
                print(
                    "poll-bytes",
                    json.dumps(
                        {
                            "before": len(json.dumps(state)),
                            "after": len(json.dumps(public)),
                        }
                    ),
                )
            self.assertEqual(
                _public_status(
                    {
                        "state": "running",
                        "operation": "build",
                        "progress": 90,
                        "output_mode": "both",
                    }
                ),
                "Building both files",
            )

    def test_scanned_pdf_prompt_preference_is_validated(self) -> None:
        settings = _validate_settings(
            {"prompt_for_scanned_pdfs": False, "setup_version": 1}
        )
        self.assertFalse(settings["prompt_for_scanned_pdfs"])
        self.assertEqual(settings["setup_version"], 1)
        self.assertEqual(settings["output_mode"], "book")
        self.assertTrue(settings["enrich_spans"])
        self.assertEqual(
            _validate_settings({"highlight_style": "paragraph"})["highlight_style"],
            "paragraph",
        )
        self.assertEqual(
            _validate_settings({"court_profile": "fca"})["court_profile"],
            "fca",
        )
        abca = _validate_settings({
            "court_profile": "abca",
            "output_mode": "book",
            "table_delivery": "native_append",
        })
        self.assertEqual(abca["output_mode"], "table")
        self.assertEqual(abca["table_delivery"], "linked_append")
        with self.assertRaises(RequestError):
            _validate_settings({"prompt_for_scanned_pdfs": "no"})
        with self.assertRaises(RequestError):
            _validate_settings({"setup_version": True})
        with self.assertRaises(RequestError):
            _validate_settings({"court_profile": "imaginary"})

    def test_docx_upload_runs_background_detector(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            document_path = root / "input.docx"
            document = Document()
            document.add_paragraph("R v Grant, 2009 SCC 32.")
            document.save(document_path)
            server = create_server(0, state_root=root / "state")
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base = f"http://127.0.0.1:{server.server_port}"
            try:
                request = urllib.request.Request(
                    f"{base}/api/jobs?filename=input.docx&split_fallback=auto",
                    data=document_path.read_bytes(),
                    headers={
                        "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    },
                    method="POST",
                )
                with urllib.request.urlopen(request) as response:
                    job = json.load(response)
                for _ in range(100):
                    with urllib.request.urlopen(f"{base}/api/jobs/{job['id']}") as response:
                        job = json.load(response)
                    if job["state"] != "running":
                        break
                    time.sleep(0.05)
                self.assertEqual(job["state"], "complete", job.get("log"))
                self.assertEqual(job["split_fallback"], "auto")
                with urllib.request.urlopen(f"{base}/api/jobs/{job['id']}/review") as response:
                    review = json.load(response)
                self.assertEqual(Path(review["input_path"]).name, "input.docx")
                self.assertGreaterEqual(len(review["parts"]), 1)
                self.assertEqual(review["split_fallback"]["requested"], "auto")
                self.assertEqual(review["split_fallback"]["eligible_units"], 0)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=2)

    def test_pdf_upload_uses_same_source_for_detection_and_book_build(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            state_root = Path(temporary) / "state"
            commands: list[tuple[str, list[str]]] = []

            def finish_immediately(
                job: Path,
                operation: str,
                argv: list[str],
                *,
                success=None,
            ) -> None:
                commands.append((operation, argv))
                if operation == "detection":
                    (job / "review.json").write_text("{}", encoding="utf-8")
                    if success:
                        success()
                _update_job(job, state="complete", operation=operation, progress=100)

            with mock.patch("toa_web._start_process", side_effect=finish_immediately):
                server = create_server(0, state_root=state_root)
                thread = threading.Thread(target=server.serve_forever, daemon=True)
                thread.start()
                base = f"http://127.0.0.1:{server.server_port}"
                try:
                    request = urllib.request.Request(
                        f"{base}/api/jobs?filename=brief.pdf",
                        data=b"%PDF-1.7\n%%EOF",
                        headers={"Content-Type": "application/pdf"},
                        method="POST",
                    )
                    with urllib.request.urlopen(request) as response:
                        job = json.load(response)

                    source = state_root / "jobs" / job["id"] / "input.pdf"
                    self.assertTrue(source.is_file())
                    self.assertEqual(job["input_file"], "input.pdf")
                    self.assertEqual(commands[0][0], "detection")
                    self.assertIn(str(source), commands[0][1])

                    request = urllib.request.Request(
                        f"{base}/api/jobs/{job['id']}/build",
                        data=json.dumps({"output_mode": "table"}).encode(),
                        headers={"Content-Type": "application/json"},
                        method="POST",
                    )
                    with self.assertRaises(urllib.error.HTTPError) as error:
                        urllib.request.urlopen(request)
                    self.assertEqual(error.exception.code, 400)
                    self.assertEqual(
                        json.load(error.exception)["error"],
                        "PDF documents can create a Book of Authorities only.",
                    )
                    self.assertEqual(len(commands), 1)

                    request = urllib.request.Request(
                        f"{base}/api/jobs/{job['id']}/build",
                        data=json.dumps({"output_mode": "book"}).encode(),
                        headers={"Content-Type": "application/json"},
                        method="POST",
                    )
                    with urllib.request.urlopen(request) as response:
                        built = json.load(response)
                    self.assertEqual(built["output_mode"], "book")
                    self.assertEqual(commands[-1][0], "build")
                    self.assertIn(str(source), commands[-1][1])
                finally:
                    server.shutdown()
                    server.server_close()
                    thread.join(timeout=2)

    def test_docx_upload_rejects_unknown_split_fallback(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            document_path = root / "input.docx"
            document = Document()
            document.add_paragraph("R v Grant, 2009 SCC 32.")
            document.save(document_path)
            server = create_server(0, state_root=root / "state")
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base = f"http://127.0.0.1:{server.server_port}"
            try:
                request = urllib.request.Request(
                    f"{base}/api/jobs?filename=input.docx&split_fallback=unbounded",
                    data=document_path.read_bytes(),
                    method="POST",
                )
                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(request)
                self.assertEqual(error.exception.code, 400)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=2)

    def test_manual_pdf_state_and_build(self) -> None:
        import fitz

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = fitz.open()
            page = source.new_page()
            page.insert_text((72, 72), "Authority")
            pdf = source.tobytes()
            source.close()
            server = create_server(0, state_root=root)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base = f"http://127.0.0.1:{server.server_port}"
            try:
                request = urllib.request.Request(f"{base}/api/jobs", data=b"", method="POST")
                with urllib.request.urlopen(request) as response:
                    job = json.load(response)
                request = urllib.request.Request(
                    f"{base}/api/jobs/{job['id']}/manual/files?filename=authority.pdf",
                    data=pdf,
                    headers={"Content-Type": "application/pdf"},
                    method="POST",
                )
                with urllib.request.urlopen(request) as response:
                    manual = json.load(response)
                manual["book_title"] = "Test Book"
                manual["entries"][0]["tab"] = "A"
                request = urllib.request.Request(
                    f"{base}/api/jobs/{job['id']}/manual/build",
                    data=json.dumps(manual).encode(),
                    headers={"Content-Type": "application/json"},
                    method="POST",
                )
                with urllib.request.urlopen(request):
                    pass
                for _ in range(100):
                    with urllib.request.urlopen(f"{base}/api/jobs/{job['id']}") as response:
                        job = json.load(response)
                    if job["state"] != "running":
                        break
                    time.sleep(0.05)
                self.assertEqual(job["state"], "complete", job.get("log"))
                self.assertTrue(any(row["name"] == "Test-Book.pdf" for row in job["files"]))
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=2)


if __name__ == "__main__":
    unittest.main()
