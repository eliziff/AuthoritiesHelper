import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from toa_maker import (
    A2AJClient,
    A2AJDocument,
    A2AJLookup,
    Authority,
    DeterministicPart,
    Occurrence,
    ReviewState,
    TextUnit,
    _authority_sourcedoc_payload,
    _authority_pdf_document,
    _download_pdf,
    _pinpoint_areas,
    _pinpoint_pattern,
    _prepare_authority_pdf,
    _merge_authority,
    _part_from,
    _render_pdf,
    _set_authority_span,
    _set_pinpoint_span,
    annotate_docx,
    attach_manifest_extra_pdf,
    attach_manual_pdf,
    apply_docx_discrepancy,
    assign_tabs,
    build_manual_book,
    build_project,
    extract_fields,
    extract_pdf_units,
    extract_text_fields,
    finalize_manifest_book,
    review_document,
    resolve_review,
    split_citations,
    write_book_of_authorities,
    write_combined_book_pdf,
    write_table_of_authorities,
    write_table_of_authorities_pdf,
)
from project_store import load_project_file, save_project_file
from quote_edits import editorial_quote


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _docx_with_footnote(path: Path, text: str) -> None:
    document = f'''<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>Text.</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p><w:sectPr/></w:body></w:document>'''
    footnotes = f'''<?xml version="1.0" encoding="UTF-8"?>
<w:footnotes xmlns:w="{W}">
<w:footnote w:id="-1" w:type="separator"><w:p/></w:footnote>
<w:footnote w:id="1"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:footnote>
</w:footnotes>'''
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr("word/footnotes.xml", footnotes)


def _docx_with_quote_and_footnote(path: Path, quote: str, citation: str) -> None:
    document = f'''<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>The Court held that “{quote}”.</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p><w:sectPr/></w:body></w:document>'''
    footnotes = f'''<?xml version="1.0" encoding="UTF-8"?>
<w:footnotes xmlns:w="{W}">
<w:footnote w:id="-1" w:type="separator"><w:p/></w:footnote>
<w:footnote w:id="1"><w:p><w:r><w:t>{citation}</w:t></w:r></w:p></w:footnote>
</w:footnotes>'''
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr("word/footnotes.xml", footnotes)


class SplitterTests(unittest.TestCase):
    def test_docx_rejects_xml_declarations(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "hostile.docx"
            with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("word/document.xml", "<!DOCTYPE x [<!ENTITY y 'z'>]><x>&y;</x>")
            with self.assertRaisesRegex(ValueError, "declarations are not allowed"):
                review_document(source)

    def test_docx_rejects_active_embedded_content(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "hostile.docx"
            with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("word/document.xml", f'<w:document xmlns:w="{W}"/>')
                archive.writestr("word/embeddings/oleObject1.bin", b"untrusted")
            with self.assertRaisesRegex(ValueError, "active or embedded content"):
                review_document(source)

    def test_pdf_notes_feed_the_same_review_units_as_word_notes(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "factum.pdf"
            with fitz.open() as document:
                page = document.new_page(width=612, height=792)
                page.insert_text(
                    (72, 90),
                    "The first proposition is supported by authority¹.",
                    fontsize=11,
                )
                page.draw_line((72, 650), (220, 650), width=0.7)
                page.insert_text(
                    (72, 675),
                    "1 R v Grant, 2009 SCC 32.",
                    fontsize=8,
                )
                document.save(source)

            with patch.dict(
                "os.environ",
                {"OPEN_LEGAL_DATA_HOME": str(root / "legal-data")},
            ):
                units = extract_pdf_units(source)
                review = review_document(source)

            self.assertEqual(
                ["R v Grant, 2009 SCC 32."],
                [unit.text for unit in units if unit.kind == "footnote"],
            )
            self.assertEqual(
                [1],
                [
                    note
                    for unit in units
                    if unit.kind == "body"
                    for note, _offset in unit.footnote_refs
                ],
            )
            self.assertEqual(
                ["2009 SCC 32"],
                [part.bare_citation for part in review.parts],
            )

    def test_top_level_semicolon_and_parallel_reporters(self):
        result = split_citations("R v Oakes, [1986] 1 SCR 103; R v Sparrow, [1990] 1 SCR 1075.")
        self.assertEqual(result.status, "deterministic_complete")
        self.assertEqual(len(result.parts), 2)
        self.assertEqual(result.parts[0].text, "R v Oakes, [1986] 1 SCR 103")

    def test_us_reporter_and_statute_use_shared_grammar_anchors(self):
        result = split_citations(
            "Roe v Wade, 410 U.S. 113; claim under 42 U.S.C. § 1983."
        )
        self.assertEqual(
            [part.anchors for part in result.parts],
            [("reporter",), ("statute",)],
        )

    def test_parenthetical_party_name_is_kept_in_full_case_span(self):
        result = split_citations(
            "Ibid at para 38. See R v Imperial Tobacco, 2011 SCC 42; "
            "Nelson (City) v Marchi, 2021 SCC 41."
        )
        self.assertEqual(
            [extract_fields(part).citation_with_style for part in result.parts],
            [
                "Ibid",
                "R v Imperial Tobacco, 2011 SCC 42",
                "Nelson (City) v Marchi, 2021 SCC 41",
            ],
        )

    def test_semicolon_in_title_is_not_boundary(self):
        result = split_citations('Jane Doe, “A Title; With a Subtitle” (2020) 1 Queen\'s LJ 10; R v Oakes, [1986] 1 SCR 103.')
        self.assertEqual(len(result.parts), 2)

    def test_optional_fallback_receives_only_incomplete_units(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "brief.docx"
            text = "R v Grant, 2009 SCC 32; explanatory residue."
            _docx_with_footnote(source, text)
            calls = []

            def planner(records, **options):
                calls.append((records, options))
                return {
                    "strategy_used": "hybrid",
                    "footnotes": [
                        {
                            **records[0],
                            "parts": [
                                {
                                    "verbatim": "R v Grant, 2009 SCC 32",
                                    "kind": "case",
                                    "bare_citation": "2009 SCC 32",
                                    "citation_with_style": "R v Grant, 2009 SCC 32",
                                    "short_form": "R v Grant",
                                    "pinpoint_fragments": [],
                                    "page_pinpoints": [],
                                    "route": "codex",
                                },
                                {
                                    "verbatim": "explanatory residue.",
                                    "kind": "other",
                                    "bare_citation": "explanatory residue",
                                    "citation_with_style": "explanatory residue.",
                                    "short_form": "",
                                    "pinpoint_fragments": [],
                                    "page_pinpoints": [],
                                    "route": "codex",
                                },
                            ],
                        }
                    ],
                    "telemetry": {
                        "codex_batches": 1,
                        "live_codex_batches": 1,
                        "token_usage": {"input_tokens": 100},
                    },
                }

            review = review_document(
                source,
                split_fallback="auto",
                split_model="test-model",
                split_effort="high",
                split_planner=planner,
            )

            self.assertEqual(len(calls), 1)
            self.assertEqual([row["text"] for row in calls[0][0]], [text])
            self.assertEqual(calls[0][1]["strategy"], "hybrid")
            self.assertEqual(
                [part.split_status for part in review.parts],
                ["codex_fallback", "codex_fallback"],
            )
            self.assertEqual(review.split_fallback["eligible_units"], 1)
            self.assertEqual(review.split_fallback["live_codex_batches"], 1)

    def test_fields_keep_kind_and_pinpoint(self):
        fields = extract_text_fields("R v X, 2020 SCC 1 at para. 20")
        self.assertEqual(fields.kind, "case")
        self.assertEqual(fields.bare_citation, "2020 SCC 1")
        self.assertEqual(fields.citation_with_style, "R v X, 2020 SCC 1")
        self.assertEqual(fields.pinpoint_fragments, ("par20",))
        law = extract_text_fields("RSC 1985, c C-46, s 16")
        self.assertEqual(law.kind, "statute")
        self.assertEqual(law.bare_citation, "RSC 1985, c C-46")
        self.assertEqual(law.pinpoint_fragments, ("sec16",))

    def test_multiple_pinpoints_are_occurrence_data(self):
        fields = extract_text_fields("R v X, 2020 SCC 1 at paras 20, 23 and 25")
        self.assertEqual(fields.citation_with_style, "R v X, 2020 SCC 1")
        self.assertEqual(fields.pinpoint_fragments, ("par20", "par23", "par25"))

    def test_exact_page_range_expands_and_malformed_range_fails_closed(self):
        fields = extract_text_fields("R v X, [1986] 1 SCR 103 at pp 104-106")
        self.assertEqual(fields.page_pinpoints, (104, 105, 106))
        self.assertEqual(
            extract_text_fields("R v X, [1986] 1 SCR 103 at pp 106-104").page_pinpoints,
            (),
        )

    def test_chamberlain_first_footnote_keeps_only_primary_authority(self):
        fields = extract_text_fields("2023 SCC 14, rev’g 2021 BCCA 222 [Hansman].")
        self.assertEqual(fields.citation_with_style, "2023 SCC 14")
        self.assertEqual(fields.bare_citation, "2023 SCC 14")
        self.assertEqual(fields.short_form, "Hansman")

    def test_subsequent_history_signals_split_with_exact_offsets(self):
        text = "2023 SCC 14, rev’g 2021 BCCA 222 [Hansman]."
        result = split_citations(text)
        self.assertEqual(result.status, "deterministic_complete")
        self.assertEqual(
            [part.text for part in result.parts],
            ["2023 SCC 14", "rev’g 2021 BCCA 222 [Hansman]."],
        )
        self.assertEqual(
            [extract_fields(part).bare_citation for part in result.parts],
            ["2023 SCC 14", "2021 BCCA 222"],
        )
        second = result.parts[1]
        self.assertEqual(second.start, text.index("rev’g"))
        self.assertEqual(second.end, len(text))
        self.assertEqual(text[second.start:second.end], second.text)

        for signal in (
            "rev'g",
            "rev’g",
            "rev'd",
            "rev’d",
            "aff'g",
            "aff’g",
            "aff'd",
            "aff’d",
            "cited by",
        ):
            with self.subTest(signal=signal):
                value = f"R v First, 2020 SCC 1 {signal} R v Second, 2021 SCC 2."
                split = split_citations(value)
                self.assertEqual(split.status, "deterministic_complete")
                self.assertEqual(split.parts[1].start, value.index(signal))
                self.assertEqual(
                    [part.text for part in split.parts],
                    ["R v First, 2020 SCC 1", f"{signal} R v Second, 2021 SCC 2."],
                )
                self.assertEqual(
                    [extract_fields(part).bare_citation for part in split.parts],
                    ["2020 SCC 1", "2021 SCC 2"],
                )

    def test_journal_author_and_title_are_part_of_the_authority(self):
        text = (
            "On related policy, see Dan Priel, “The Political Origins of English Private Law” "
            "(2013) 40:4 J Law and Society 481."
        )
        fields = extract_text_fields(text)
        self.assertEqual(
            fields.citation_with_style,
            "Dan Priel, “The Political Origins of English Private Law” (2013) 40:4 J Law and Society 481",
        )
        self.assertEqual(fields.bare_citation, "(2013) 40:4 J Law and Society 481")

    def test_statute_drops_signal_and_section_pinpoint(self):
        fields = extract_text_fields(
            "The preamble states several purposes. See Courts of Justice Act, RSO 1990, c C.43, s 137."
        )
        self.assertEqual(fields.citation_with_style, "Courts of Justice Act, RSO 1990, c C.43")
        self.assertEqual(fields.bare_citation, "RSO 1990, c C.43")
        self.assertEqual(fields.pinpoint_fragments, ("sec137",))

    def test_large_surface_keeps_every_boundary(self):
        text = "; ".join(f"R v Example {index}, 2020 SCC {index}" for index in range(1, 501))
        result = split_citations(text)
        self.assertEqual(result.status, "deterministic_complete")
        self.assertEqual(len(result.parts), 500)


class BuildTests(unittest.TestCase):
    def test_pdf_source_refuses_word_table_output(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "factum.pdf"
            with fitz.open() as document:
                document.new_page()
                document.save(source)

            with self.assertRaisesRegex(ValueError, "only into a Word document"):
                build_project(
                    source,
                    root / "output",
                    output_mode="table",
                    offline=True,
                    pdf_mode="none",
                )

    def test_default_tabs_are_grouped_and_alphabetical(self):
        authorities = [
            Authority("z-case", "case", "2020 SCC 2", "Zulu v Alpha"),
            Authority("a-statute", "statute", "RSC 1985, c A-1", "Access Act"),
            Authority("a-case", "case", "2020 SCC 1", "Alpha v Zulu"),
            Authority("a-journal", "journal", "(2020) 1 Law Review 1", "A Legal Article"),
        ]

        assign_tabs(authorities)

        self.assertEqual(
            [(authority.kind, authority.name, authority.tab) for authority in authorities],
            [
                ("case", "Alpha v Zulu", "Tab 1"),
                ("case", "Zulu v Alpha", "Tab 2"),
                ("statute", "Access Act", "Tab 3"),
                ("journal", "A Legal Article", "Tab 4"),
            ],
        )

    def test_pdf_source_modes_have_distinct_fallbacks(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            cases = (
                ("auto", "reconstructed", "Available source text"),
                ("originals", "placeholder", "SOURCE PDF UNAVAILABLE"),
                ("render", "reconstructed", "Available source text"),
            )
            for mode, origin, expected_text in cases:
                with self.subTest(mode=mode):
                    authority = Authority(
                        mode,
                        "case",
                        "2020 SCC 1",
                        "Example v Example",
                        source_text="Available source text",
                        tab="Tab 1",
                    )
                    _download_pdf(authority, root / mode, mode)
                    self.assertEqual(authority.pdf_origin, origin)
                    with fitz.open(authority.pdf_path) as document:
                        self.assertIn(expected_text, "\n".join(page.get_text() for page in document))

    def test_reconstructed_pdf_uses_canonical_slices_or_honest_flat_text(self):
        import fitz

        canonical = {
            "kind": "sections",
            "segments": [
                {
                    "kind": "section",
                    "label": "sec9",
                    "origin": "native",
                    "text": "Provider-map wording.",
                }
            ],
        }
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            authority = Authority(
                "law",
                "statute",
                "Test Act",
                "Test Act",
                source_text="Whole-text wording that must not be rendered.",
                tab="Tab 1",
            )
            with patch(
                "toa_maker._compile_authority_sourcedoc",
                return_value=canonical,
            ):
                _download_pdf(authority, root / "canonical", "render")
            with fitz.open(authority.pdf_path) as document:
                text = "\n".join(page.get_text() for page in document)
            self.assertIn("Section 9", text)
            self.assertIn("Provider-map wording.", text)
            self.assertNotIn("Whole-text wording", text)

            flat = Authority(
                "case",
                "case",
                "2020 SCC 1",
                "Example v Example",
                source_text="[1] Flat source text remains exactly visible.",
                tab="Tab 2",
            )
            with patch(
                "toa_maker._compile_authority_sourcedoc",
                return_value=None,
            ):
                _download_pdf(flat, root / "flat", "render")
            with fitz.open(flat.pdf_path) as document:
                text = "\n".join(page.get_text() for page in document)
            self.assertIn("[1] Flat source text remains exactly visible.", text)

    @staticmethod
    def _image_only_pdf(path: Path, text: str, pages: int = 1) -> None:
        import fitz

        source = fitz.open()
        page = source.new_page(width=612, height=792)
        page.insert_text((72, 100), text, fontsize=16)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        source.close()
        scanned = fitz.open()
        for _ in range(pages):
            target = scanned.new_page(width=612, height=792)
            target.insert_image(target.rect, stream=pixmap.tobytes("png"))
        scanned.save(path)
        scanned.close()

    def test_scanned_pdf_uses_canonical_ocr_and_preserves_original(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "scan.pdf"
            self._image_only_pdf(source, "[20] The exact quotation appears here.")
            original = source.read_bytes()
            authority = Authority(
                "scan",
                "case",
                "2020 SCC 1",
                "Example v Example",
                pdf_path=str(source),
                pdf_origin="original",
                occurrences=[
                    Occurrence(
                        "p1", "footnote:1", "case", 1, 1, "", "2020 SCC 1",
                        ["par20"], [], exact_quotes=["The exact quotation appears here."],
                    )
                ],
            )
            _prepare_authority_pdf(authority, "cited_pages")
            self.assertFalse(authority.pdf_has_text_layer)
            self.assertEqual(authority.ocr_scope, "cited_pages")
            self.assertEqual(authority.ocr_pages, [0])
            self.assertEqual(source.read_bytes(), original)
            self.assertIsNotNone(getattr(authority, "_pdf_document", None))
            evidence = _authority_pdf_document(authority)
            self.assertIn(
                "exact quotation",
                " ".join(line.text for line in evidence.pages[0].lines).casefold(),
            )

    def test_full_ocr_processes_every_scanned_page(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "scan.pdf"
            self._image_only_pdf(source, "A scanned authority", pages=2)
            authority = Authority(
                "scan",
                "case",
                "2020 SCC 1",
                "Example v Example",
                pdf_path=str(source),
                pdf_origin="original",
                occurrences=[
                    Occurrence("p1", "footnote:1", "case", 1, 1, "", "", ["par20"], [])
                ],
            )
            _prepare_authority_pdf(authority, "full")
            self.assertEqual(authority.ocr_scope, "full")
            self.assertEqual(authority.ocr_pages, [0, 1])

    def test_page_pinpoint_without_an_artifact_label_fails_closed(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "scan.pdf"
            output = root / "book.pdf"
            self._image_only_pdf(source, "A scanned authority", pages=2)
            authority = Authority(
                "scan",
                "case",
                "[1986] 1 SCR 103",
                "Example v Example",
                tab="Tab 1",
                pdf_path=str(source),
                pdf_origin="original",
                pdf_has_text_layer=False,
                occurrences=[
                    Occurrence("p1", "footnote:1", "case", 1, 1, "", "", [], [103])
                ],
            )
            write_combined_book_pdf([authority], output, "brief.docx", highlight_style="margin")
            with fitz.open(output) as document:
                first_authority_page = document[2]
                bars = [
                    drawing["rect"]
                    for drawing in first_authority_page.get_drawings()
                    if drawing["rect"].width <= 7 and drawing["rect"].height > 700
                ]
                self.assertEqual(bars, [])

    def test_page_pinpoint_uses_the_engine_printed_page_label(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "labelled.pdf"
            output = root / "book.pdf"
            with fitz.open() as document:
                for label in (103, 104):
                    page = document.new_page(width=612, height=792)
                    page.insert_text((72, 100), f"Authority text on page {label}.")
                    page.insert_text((298, 760), str(label))
                document.save(source)
            authority = Authority(
                "labelled",
                "case",
                "[1986] 1 SCR 103",
                "Example v Example",
                tab="Tab 1",
                pdf_path=str(source),
                pdf_origin="manual",
                occurrences=[
                    Occurrence("p1", "footnote:1", "case", 1, 1, "", "", [], [104])
                ],
            )

            write_combined_book_pdf(
                [authority], output, "brief.docx", highlight_style="margin"
            )

            with fitz.open(output) as document:
                bars = [
                    [
                        drawing["rect"]
                        for drawing in document[index].get_drawings()
                        if drawing.get("fill") and drawing["rect"].width <= 7
                    ]
                    for index in (2, 3)
                ]
            self.assertEqual(bars[0], [])
            self.assertEqual(len(bars[1]), 1)

    def test_combined_book_reuses_the_prepared_pdf_document(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source.pdf"
            output = root / "book.pdf"
            with fitz.open() as document:
                document.new_page().insert_text(
                    (72, 100), "[20] The cited paragraph appears here."
                )
                document.save(source)
            authority = Authority(
                "cached",
                "case",
                "2020 SCC 1",
                "Example v Example",
                tab="Tab 1",
                pdf_path=str(source),
                pdf_origin="manual",
                occurrences=[
                    Occurrence(
                        "p1", "footnote:1", "case", 1, 1, "", "", ["par20"], []
                    )
                ],
            )
            with patch("toa_maker._tesseract_command", return_value=None):
                _prepare_authority_pdf(authority, "cited_pages")

            with patch(
                "toa_maker._prepare_authority_pdf",
                side_effect=AssertionError("authority PDF was reparsed"),
            ):
                write_combined_book_pdf(
                    [authority], output, "brief.docx", highlight_style="margin"
                )

    def test_section_pinpoint_uses_the_engine_section_span(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "law.pdf"
            with fitz.open() as document:
                page = document.new_page(width=612, height=792)
                page.insert_text((72, 100), "137 Definitions", fontsize=16)
                page.insert_text(
                    (72, 135), "This section applies to the defined term.", fontsize=11
                )
                page.insert_text((72, 170), "138 Application", fontsize=16)
                page.insert_text(
                    (72, 205), "The next section begins here.", fontsize=11
                )
                document.save(source)
            authority = Authority(
                "law",
                "statute",
                "RSC 1985, c E-1",
                "Example Act",
                pdf_path=str(source),
                pdf_origin="manual",
                occurrences=[
                    Occurrence(
                        "p1", "footnote:1", "statute", 1, 1, "", "", ["sec137"], []
                    )
                ],
            )
            document = _prepare_authority_pdf(authority, "page_margin")
            pattern = _pinpoint_pattern("sec137")
            self.assertIsNotNone(pattern)
            areas = _pinpoint_areas(document, "sec137", pattern)

            self.assertEqual(list(areas), [0])
            self.assertEqual(len(areas[0]), 1)
            self.assertLess(areas[0][0][1], 100)
            self.assertGreater(areas[0][0][3], 135)
            self.assertLess(areas[0][0][3], 170)

    def test_production_pdf_paths_do_not_extract_text_with_fitz(self):
        root = Path(__file__).parents[1]
        source = "\n".join(
            path.read_text(encoding="utf-8") for path in root.glob("*.py")
        )
        for private_extractor in (
            ".get_text(",
            ".get_textpage_ocr(",
            ".search_for(",
        ):
            self.assertNotIn(private_extractor, source)

    def test_wrong_pinpoint_never_highlights_quote(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "native.pdf"
            output = root / "book.pdf"
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 90), "[20] The court found this carefully stated proposition was correct.")
            document.save(source)
            document.close()
            authority = Authority(
                "native",
                "case",
                "2020 SCC 1",
                "Example v Example",
                tab="Tab 1",
                pdf_path=str(source),
                pdf_origin="original",
                pdf_has_text_layer=True,
                occurrences=[
                    Occurrence(
                        "p1", "footnote:1", "case", 1, 1, "", "", ["par19"], [],
                        exact_quotes=["The court found this carefully stated proposition is correct."],
                        pinpoint_text="at para. 19",
                        pinpoint_start=24,
                        pinpoint_end=35,
                    )
                ],
            )
            discrepancies = write_combined_book_pdf(
                [authority],
                output,
                "brief.docx",
                highlight_style="text",
            )
            self.assertEqual(discrepancies, [])
            with fitz.open(output) as combined:
                page = combined[2]
                self.assertEqual(list(page.annots() or []), [])

    def test_source_pinpoint_changes_only_after_explicit_acceptance_and_creates_backup(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            text = "R v Example, 2020 SCC 1 at para. 19."
            _docx_with_footnote(source, text)
            start = text.index("at para. 19")
            manifest = root / "brief.toa-manifest.json"
            manifest.write_text(
                json.dumps(
                    {
                        "input": str(source),
                        "discrepancies": [
                            {
                                "id": "change-1",
                                "status": "pending",
                                "unit_key": "footnote:1",
                                "source_pinpoint": "at para. 19",
                                "pinpoint_start": start,
                                "pinpoint_end": start + len("at para. 19"),
                                "replacement_text": "at para. 20",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            before = source.read_bytes()
            backup = apply_docx_discrepancy(
                manifest,
                "change-1",
                action="pinpoint",
            )
            self.assertIsNotNone(backup)
            self.assertEqual(Path(backup).read_bytes(), before)
            with zipfile.ZipFile(source) as archive:
                footnotes = archive.read("word/footnotes.xml").decode("utf-8")
            self.assertIn("at para. 20", footnotes)
            updated = json.loads(manifest.read_text(encoding="utf-8"))
            self.assertEqual(updated["discrepancies"][0]["status"], "applied")

    def test_quote_discrepancy_can_apply_exact_source_wording(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            _docx_with_quote_and_footnote(
                source,
                "This and that",
                "Example v Example, 2020 SCC 1 at para 20.",
            )
            review = review_document(source)
            manifest = root / "brief.toa-manifest.json"
            manifest.write_text(
                json.dumps(
                    {
                        "input": str(source),
                        "review": review.to_dict(),
                        "discrepancies": [
                            {
                                "id": "quote-1",
                                "status": "pending",
                                "reasons": ["quote_text"],
                                "footnote_id": 1,
                                "expected_quote": "This and that",
                                "found_quote": "This long passage and another",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            backup = apply_docx_discrepancy(manifest, "quote-1", action="quote_exact")
            self.assertTrue(Path(backup).is_file())
            with zipfile.ZipFile(source) as archive:
                document = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("This long passage and another", document)
            updated = json.loads(manifest.read_text(encoding="utf-8"))
            self.assertEqual(updated["discrepancies"][0]["action"], "quote_exact")

    def test_editorial_quote_uses_source_wording_with_brackets_and_ellipsis(self):
        self.assertEqual(
            editorial_quote("This and that", "This long passage and another"),
            "This ... and [that]",
        )

    def test_lightweight_project_file_is_versioned_and_atomic(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "matter.toa-project.json"
            save_project_file(path, {"input_path": "brief.docx", "review": {"parts": []}})
            self.assertFalse(path.with_name(f"{path.name}.part").exists())
            loaded = load_project_file(path)
            self.assertEqual(loaded["format"], "table-of-authorities-project")
            self.assertEqual(loaded["version"], 1)
            self.assertEqual(loaded["input_path"], "brief.docx")

    def test_a2aj_metadata_is_not_mistaken_for_source_text(self):
        self.assertEqual(
            A2AJClient._text(
                {
                    "dataset": "SCC",
                    "citation_en": "2099 SCC 999",
                    "name_en": "Example v. Missing",
                    "document_date_en": "2099-01-01",
                }
            ),
            "",
        )

    def test_reporter_lookup_uses_case_name_and_preserves_offline_search(self):
        class Client(A2AJClient):
            def __init__(self, searched):
                self.offline = False
                self.searched = searched
                self.calls = []

            def _get(self, endpoint, params):
                self.calls.append((endpoint, params))
                if endpoint == "/fetch":
                    return {"status": 200, "json": {}, "error": ""}
                return self.searched

        paice = {
            "dataset": "SCC",
            "citation_en": "2005 SCC 22",
            "citation2_en": "[2005] 1 SCR 339",
            "name_en": "R. v. Paice",
            "source_url_en": "https://decisions.scc-csc.ca/scc-csc/scc-csc/en/item/2222/index.do",
        }
        client = Client({"status": 200, "json": {"results": [paice]}, "error": ""})
        lookup = client.lookup(
            "[2005] 1 SCR 339",
            "case",
            name="R v Paice",
        )
        self.assertEqual(lookup.status, "found")
        self.assertEqual(client.calls[1][1]["query"], "R v Paice")
        self.assertEqual(client.calls[1][1]["search_type"], "name")

        offline = Client({"status": None, "json": None, "error": "offline"})
        lookup = offline.lookup("[2005] 1 SCR 339", "case")
        self.assertEqual(lookup.status, "offline")
        self.assertEqual(lookup.method, "exact_search")

    def test_bare_reporter_resolves_to_full_a2aj_style_of_cause(self):
        class Client:
            offline = False

            def lookup(self, citation, kind, **_):
                self.request = (citation, kind)
                return A2AJLookup(
                    "found",
                    A2AJDocument(
                        "SCC",
                        "[1986] 1 SCR 103",
                        "1986 CanLII 46 (SCC)",
                        "R. v. Oakes",
                        "1986-02-28",
                        "https://example.test/oakes",
                        "decision text",
                        "en",
                    ),
                    "exact_citation",
                )

        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "brief.docx"
            _docx_with_footnote(source, "[1986] 1 SCR 103.")
            review = review_document(source)
            self.assertEqual(review.parts[0].short_form, "")
            client = Client()
            result = resolve_review(review, client)
            self.assertEqual(client.request, ("[1986] 1 SCR 103", "case"))
            self.assertEqual(result.authorities[0].name, "R. v. Oakes")

    def test_bare_legislation_resolves_to_full_a2aj_name(self):
        class Client:
            offline = False

            def lookup(self, citation, kind, **_):
                self.request = (citation, kind)
                return A2AJLookup(
                    "found",
                    A2AJDocument(
                        "LEGISLATION-ON",
                        "RSO 1990, c C.43",
                        "",
                        "Courts of Justice Act",
                        "2025-12-11",
                        "https://example.test/courts-of-justice-act",
                        "statute text",
                        "en",
                        {
                            "unofficial_sections_en": {
                                "137": "Provider section text"
                            }
                        },
                    ),
                    "exact_citation",
                )

        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "brief.docx"
            _docx_with_footnote(source, "RSO 1990, c C.43, s 137.")
            review = review_document(source)
            client = Client()
            result = resolve_review(review, client)
            self.assertEqual(client.request, ("RSO 1990, c C.43", "statute"))
            self.assertEqual(result.authorities[0].name, "Courts of Justice Act")
            self.assertEqual(result.authorities[0].occurrences[0].pinpoint_fragments, ["sec137"])
            self.assertEqual(
                result.authorities[0].source_sections,
                {"137": "Provider section text"},
            )
            self.assertEqual(
                _authority_sourcedoc_payload(result.authorities[0])["sectionMap"],
                {"137": "Provider section text"},
            )

    def test_resolved_name_replaces_provisional_name_when_aliases_merge(self):
        provisional = Authority("a", "case", "[1986] 1 SCR 103", "[1986]", lookup_status="not_found")
        resolved = Authority(
            "b",
            "case",
            "[1986] 1 SCR 103",
            "R. v. Oakes",
            lookup_status="found",
        )
        merged = _merge_authority(provisional, resolved)
        self.assertEqual(merged.name, "R. v. Oakes")

    def test_reconstruction_keeps_real_tables_emphasis_and_nested_lists(self):
        import fitz

        source = (
            "# Example Act\n\n"
            "### Table\n\n"
            "| Column 1 | Column 2 |\n"
            "| --- | --- |\n"
            "| **Former** name | *Current* name |\n"
            "| Old Court | New Court |\n\n"
            "- first bullet\n"
            "  - nested bullet\n"
            "1. first ordered item\n"
            "  1) nested ordered item\n\n"
            "6 Main provision\n"
            "(a) child provision\n"
            "(i) grandchild provision"
        )
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "structured.pdf"
            _render_pdf(output, "Tab 1", "SBC 2019, c 3", "Example Act", source)
            with fitz.open(output) as document:
                page = next(page for page in document if "Column 1" in page.get_text())
                tables = page.find_tables().tables
                self.assertTrue(any(table.row_count == 3 and table.col_count == 2 for table in tables))
                spans = [
                    span
                    for block in page.get_text("dict")["blocks"]
                    for line in block.get("lines", [])
                    for span in line["spans"]
                ]
                self.assertTrue(any("Former" in span["text"] and "Bold" in span["font"] for span in spans))
                self.assertTrue(any("Current" in span["text"] and "Italic" in span["font"] for span in spans))
                for span in spans:
                    red, green, blue = fitz.sRGB_to_rgb(span["color"])
                    self.assertEqual(red, green)
                    self.assertEqual(green, blue)
                text = page.get_text().replace("ﬁ", "fi")
                self.assertIn("first bullet", text)
                self.assertIn("nested ordered item", text)

    def test_rendered_reconstruction_does_not_leak_markdown_syntax(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "law.pdf"
            _render_pdf(
                output,
                "Tab 1",
                "SBC 2019, c 3",
                "Protection of Public Participation Act",
                "### No amendments unless permitted\n\n6 *Unless* the court orders otherwise.",
            )
            with fitz.open(output) as document:
                text = "\n".join(page.get_text() for page in document)
                self.assertIn("No amendments unless permitted", text)
                self.assertNotIn("###", text)
                self.assertNotIn("*Unless*", text)

    def test_body_quote_is_attached_to_its_footnote_occurrence(self):
        class Client:
            offline = False

            def lookup(self, _citation, _kind, **_):
                return A2AJLookup(
                    "found",
                    A2AJDocument(
                        "SCC",
                        "2020 SCC 1",
                        "",
                        "Example v. Example",
                        "2020-01-01",
                        "https://example.test",
                        "[20] Freedom matters in this passage.",
                        "en",
                    ),
                    "exact_citation",
                )

        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "brief.docx"
            _docx_with_quote_and_footnote(source, "Freedom matters", "Example v Example, 2020 SCC 1 at para 20.")
            result = resolve_review(review_document(source), Client())
            occurrence = result.authorities[0].occurrences[0]
            self.assertIn("The Court held that", occurrence.proposition_text)
            self.assertEqual(occurrence.exact_quotes, ["Freedom matters"])

    def test_combined_book_highlights_exact_quote_on_text(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source_path = root / "source.pdf"
            source = fitz.open()
            page = source.new_page()
            page.insert_text((72, 100), "[20] The court held that freedom matters in this passage.")
            source.save(source_path)
            source.close()
            authority = Authority(
                "example",
                "case",
                "2020 SCC 1",
                "Example v. Example",
                tab="Tab 1",
                pdf_path=str(source_path),
                pdf_origin="manual",
                occurrences=[
                    Occurrence(
                        "part",
                        "footnote:1",
                        "case",
                        1,
                        1,
                        "2020 SCC 1 at para 20",
                        "2020 SCC 1",
                        ["par20"],
                        [],
                        exact_quotes=["freedom matters"],
                    )
                ],
            )
            output = root / "book.pdf"
            write_combined_book_pdf([authority], output, "brief.docx", highlight_style="text")
            with fitz.open(output) as document:
                page = document[-1]
                annotations = list(page.annots() or [])
                self.assertTrue(annotations)
                self.assertEqual(annotations[0].type[1], "Highlight")

    def test_combined_book_can_mark_cited_paragraph_in_margin(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source_path = root / "source.pdf"
            source = fitz.open()
            source.new_page().insert_text((72, 100), "[20] The cited paragraph appears here.")
            source.save(source_path)
            source.close()
            authority = Authority(
                "example",
                "case",
                "2020 SCC 1",
                "Example v. Example",
                tab="Tab 1",
                pdf_path=str(source_path),
                pdf_origin="manual",
                occurrences=[
                    Occurrence(
                        "part",
                        "footnote:1",
                        "case",
                        1,
                        1,
                        "2020 SCC 1 at para 20",
                        "2020 SCC 1",
                        ["par20"],
                        [],
                        exact_quotes=["cited paragraph"],
                    )
                ],
            )
            output = root / "book.pdf"
            write_combined_book_pdf([authority], output, "brief.docx", highlight_style="margin")
            with fitz.open(output) as document:
                page = document[-1]
                bars = [
                    drawing
                    for drawing in page.get_drawings()
                    if drawing.get("fill") and drawing["rect"].width <= 6
                ]
                self.assertTrue(bars)
                annotations = list(page.annots() or [])
                self.assertEqual([annotation.type[1] for annotation in annotations], ["Highlight"])
                stroke = annotations[0].colors["stroke"]
                self.assertAlmostEqual(stroke[0], stroke[1], places=4)
                self.assertAlmostEqual(stroke[1], stroke[2], places=4)

                marker = bars[0]["rect"]
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                pixel = pixmap.pixel(
                    round((marker.x0 + marker.x1) * pixmap.width / page.rect.width / 2),
                    round((marker.y0 + marker.y1) * pixmap.height / page.rect.height / 2),
                )
                self.assertEqual(pixel[0], pixel[1])
                self.assertEqual(pixel[1], pixel[2])
                self.assertLess(pixel[0], 250)

    def test_whole_paragraph_highlight_requires_explicit_style(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source_path = root / "source.pdf"
            source = fitz.open()
            source.new_page().insert_text((72, 100), "[20] The cited paragraph appears here.")
            source.save(source_path)
            source.close()
            authority = Authority(
                "example",
                "case",
                "2020 SCC 1",
                "Example v. Example",
                tab="Tab 1",
                pdf_path=str(source_path),
                pdf_origin="manual",
                occurrences=[
                    Occurrence(
                        "part",
                        "footnote:1",
                        "case",
                        1,
                        1,
                        "2020 SCC 1 at para 20",
                        "2020 SCC 1",
                        ["par20"],
                        [],
                    )
                ],
            )
            output = root / "book.pdf"
            write_combined_book_pdf([authority], output, "brief.docx", highlight_style="paragraph")
            with fitz.open(output) as document:
                page = document[-1]
                annotations = list(page.annots() or [])
                self.assertEqual([annotation.type[1] for annotation in annotations], ["Highlight"])
                self.assertFalse(
                    any(
                        drawing.get("fill") and drawing["rect"].width <= 7
                        for drawing in page.get_drawings()
                    )
                )

    def test_inexact_quote_gets_paragraph_margin_but_no_text_highlight(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source_path = root / "source.pdf"
            source = fitz.open()
            source.new_page().insert_text(
                (72, 100),
                "[20] The court found this carefully stated proposition was correct.",
            )
            source.save(source_path)
            source.close()
            authority = Authority(
                "example",
                "case",
                "2020 SCC 1",
                "Example v. Example",
                tab="Tab 1",
                pdf_path=str(source_path),
                pdf_origin="manual",
                occurrences=[
                    Occurrence(
                        "part",
                        "footnote:1",
                        "case",
                        1,
                        1,
                        "2020 SCC 1 at para 20",
                        "2020 SCC 1",
                        ["par20"],
                        [],
                        exact_quotes=[
                            "The court found this carefully stated proposition is correct."
                        ],
                    )
                ],
            )
            output = root / "book.pdf"
            discrepancies = write_combined_book_pdf(
                [authority],
                output,
                "brief.docx",
                highlight_style="margin",
            )
            self.assertTrue(discrepancies)
            with fitz.open(output) as document:
                page = document[-1]
                self.assertEqual(list(page.annots() or []), [])
                self.assertTrue(
                    any(
                        drawing.get("fill") and drawing["rect"].width <= 7
                        for drawing in page.get_drawings()
                    )
                )

    def test_margin_mark_is_right_of_text_on_rotated_cropped_page(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source_path = root / "source.pdf"
            source = fitz.open()
            page = source.new_page(width=612, height=792)
            page.insert_text((110, 120), "[20] The cited paragraph appears here.", fontsize=12)
            page.set_cropbox(fitz.Rect(50, 40, 570, 750))
            page.set_rotation(90)
            source.save(source_path)
            source.close()
            authority = Authority(
                "example",
                "case",
                "2020 SCC 1",
                "Example v. Example",
                tab="Tab 1",
                pdf_path=str(source_path),
                pdf_origin="manual",
                pdf_has_text_layer=True,
                occurrences=[
                    Occurrence(
                        "part",
                        "footnote:1",
                        "case",
                        1,
                        1,
                        "",
                        "2020 SCC 1",
                        ["par20"],
                        [],
                    )
                ],
            )
            output = root / "book.pdf"
            write_combined_book_pdf([authority], output, "brief.docx", highlight_style="margin")
            with fitz.open(output) as document:
                page = document[-1]
                self.assertEqual(page.rotation, 90)
                text = [
                    fitz.Rect(block[:4]) * page.rotation_matrix
                    for block in page.get_text("blocks")
                    if str(block[4] or "").strip()
                ]
                bars = [
                    drawing["rect"] * page.rotation_matrix
                    for drawing in page.get_drawings()
                    if drawing.get("fill")
                    and (drawing["rect"] * page.rotation_matrix).width <= 7
                ]
                self.assertEqual(len(bars), 1)
                bar = bars[0]
                adjacent = [
                    rect
                    for rect in text
                    if rect.y0 < bar.y1 and rect.y1 > bar.y0
                ]
                self.assertTrue(adjacent)
                self.assertGreater(bar.x0, max(rect.x1 for rect in adjacent))
                self.assertFalse(any(bar.intersects(rect) for rect in text))
                self.assertTrue(page.rect.contains(bar))

    def test_missing_malformed_or_unresolved_pinpoint_never_highlights(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source_path = root / "source.pdf"
            source = fitz.open()
            source.new_page().insert_text(
                (72, 100),
                "[20] The exact quotation appears here.",
            )
            source.save(source_path)
            source.close()
            cases = (
                ([], [], "missing"),
                (["parbad"], [], "malformed"),
                (["par19"], [], "unresolved paragraph"),
                ([], [99], "unresolved page"),
            )
            for fragments, pages, label in cases:
                with self.subTest(label):
                    authority = Authority(
                        label,
                        "case",
                        "2020 SCC 1",
                        "Example v. Example",
                        tab="Tab 1",
                        pdf_path=str(source_path),
                        pdf_origin="manual",
                        pdf_has_text_layer=True,
                        occurrences=[
                            Occurrence(
                                "part",
                                "footnote:1",
                                "case",
                                1,
                                1,
                                "",
                                "2020 SCC 1",
                                fragments,
                                pages,
                                exact_quotes=["The exact quotation appears here."],
                            )
                        ],
                    )
                    output = root / f"{label.replace(' ', '-')}.pdf"
                    write_combined_book_pdf(
                        [authority],
                        output,
                        "brief.docx",
                        highlight_style="margin",
                    )
                    with fitz.open(output) as document:
                        page = document[-1]
                        self.assertEqual(list(page.annots() or []), [])
                        self.assertFalse(
                            any(
                                drawing.get("fill") and drawing["rect"].width <= 7
                                for drawing in page.get_drawings()
                            )
                        )

    def test_manual_book_keeps_user_titles_tabs_and_order(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            entries = []
            for number, text in enumerate(("FIRST PDF", "SECOND PDF"), 1):
                path = root / f"{number}.pdf"
                source = fitz.open()
                source.new_page().insert_text((72, 72), text)
                source.save(path)
                source.close()
                entries.append({"pdf_path": str(path), "title": f"Document {number}", "tab": f"Exhibit {chr(64 + number)}"})
            output = build_manual_book(entries, root / "evidence.pdf")
            with fitz.open(output) as document:
                labels = [row[1] for row in document.get_toc()]
                self.assertIn("Exhibit A — Document 1", labels)
                self.assertIn("Exhibit B — Document 2", labels)
                text = "\n".join(page.get_text() for page in document)
                self.assertLess(text.index("FIRST PDF"), text.index("SECOND PDF"))

    def test_custom_exclusion_keeps_authority_in_table_but_out_of_book_manifest(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            output = root / "out"
            _docx_with_footnote(
                source,
                "R v Oakes, [1986] 1 SCR 103; R v Sparrow, [1990] 1 SCR 1075.",
            )
            result = build_project(
                source,
                output,
                offline=True,
                pdf_mode="none",
                excluded_authorities=["R v Oakes"],
            )
            excluded = next(
                authority for authority in result.authorities
                if "Oakes" in authority.name
            )
            self.assertEqual(excluded.tab, "Not reproduced")
            manifest = json.loads((output / "brief.toa-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(len(manifest["authorities"]), 1)
            self.assertIn("Sparrow", manifest["authorities"][0]["name"])

    def test_book_removes_xml_illegal_source_controls(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "book.docx"
            authority = Authority("key", "case", "2020 SCC 1", "R v Example", source_text="Valid\x0btext", tab="Tab 1")
            write_book_of_authorities([authority], output, "brief.docx")
            self.assertTrue(output.is_file())

    def test_review_json_round_trip_and_offline_build(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            _docx_with_footnote(source, "R v Oakes, [1986] 1 SCR 103; R v Sparrow, [1990] 1 SCR 1075.")
            review = review_document(source)
            self.assertEqual(len(review.parts), 2)
            review_path = root / "brief.toa-review.json"
            review.save(review_path)
            loaded = ReviewState.load(review_path)
            self.assertEqual([part.text for part in loaded.parts], [part.text for part in review.parts])
            result = build_project(
                source,
                root / "out",
                review=loaded,
                offline=True,
                pdf_mode="render",
                output_mode="both",
                highlight_style="paragraph",
            )
            self.assertEqual(len(result.authorities), 2)
            output = root / "out"
            self.assertTrue((output / "brief.annotated.docx").exists())
            self.assertTrue((output / "brief.book-of-authorities.pdf").exists())
            self.assertFalse((output / "brief.table-of-authorities.docx").exists())
            self.assertFalse((output / "brief.book-of-authorities.docx").exists())
            with zipfile.ZipFile(output / "brief.annotated.docx") as archive:
                footnotes = archive.read("word/footnotes.xml").decode("utf-8")
                document = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("[Tab 1]", footnotes)
            self.assertIn("[Tab 2]", footnotes)
            self.assertIn(r"TA \l", footnotes)
            self.assertIn(r"TOA \h", document)
            self.assertIn("R v Sparrow, [1990] 1 SCR 1075", footnotes)
            self.assertNotIn("R v Sparrow, R v Sparrow", footnotes)
            manifest = json.loads((output / "brief.toa-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(len(manifest["authorities"]), 2)
            self.assertEqual(manifest["highlight_style"], "paragraph")
            self.assertNotIn("source_text", manifest["authorities"][0])

    def test_book_and_table_output_modes_do_not_create_unrequested_artifacts(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            _docx_with_footnote(source, "R v Oakes, [1986] 1 SCR 103.")
            table_output = root / "table"
            build_project(
                source,
                table_output,
                offline=True,
                pdf_mode="none",
                output_mode="table",
                table_delivery="pdf_append",
            )
            table_manifest = json.loads(
                (table_output / "brief.toa-manifest.json").read_text(encoding="utf-8")
            )
            self.assertTrue((table_output / "brief.annotated.docx").is_file())
            self.assertFalse((table_output / "brief.table-of-authorities.docx").exists())
            self.assertFalse((table_output / "brief.document.pdf").exists())
            self.assertFalse(
                (table_output / "brief.document-with-table-of-authorities.pdf").exists()
            )
            self.assertFalse((table_output / "brief.book-of-authorities.pdf").exists())
            self.assertEqual(table_manifest["table_delivery"], "native_append")
            self.assertEqual(
                [
                    key
                    for key, value in table_manifest["outputs"].items()
                    if value
                ],
                ["annotated_docx"],
            )

            book_output = root / "book"
            build_project(
                source,
                book_output,
                offline=True,
                pdf_mode="render",
            )
            book_manifest = json.loads(
                (book_output / "brief.toa-manifest.json").read_text(encoding="utf-8")
            )
            self.assertTrue((book_output / "brief.book-of-authorities.pdf").is_file())
            self.assertFalse((book_output / "brief.annotated.docx").exists())
            self.assertFalse((book_output / "brief.book-of-authorities.docx").exists())
            self.assertEqual(
                [
                    key
                    for key, value in book_manifest["outputs"].items()
                    if value
                ],
                ["book_of_authorities_pdf"],
            )

            both_output = root / "both"
            build_project(
                source,
                both_output,
                offline=True,
                pdf_mode="render",
                output_mode="both",
            )
            both_manifest = json.loads(
                (both_output / "brief.toa-manifest.json").read_text(encoding="utf-8")
            )
            self.assertTrue((both_output / "brief.annotated.docx").is_file())
            self.assertTrue((both_output / "brief.book-of-authorities.pdf").is_file())
            self.assertEqual(
                [
                    key
                    for key, value in both_manifest["outputs"].items()
                    if value
                ],
                ["annotated_docx", "book_of_authorities_pdf"],
            )

    def test_linked_table_is_appended_in_first_reference_order(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            output = root / "linked.docx"
            _docx_with_footnote(source, "R v Grant, 2009 SCC 32.")
            annotate_docx(
                source,
                output,
                {},
                linked_toa=[
                    ("R v Grant, 2009 SCC 32", "https://example.test/grant"),
                    ("Old text", ""),
                ],
            )
            with zipfile.ZipFile(output) as archive:
                document = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("Table of Authorities", document)
            self.assertLess(document.index("R v Grant"), document.index("Old text"))
            self.assertIn('HYPERLINK "https://example.test/grant"', document)
            self.assertIn("copy required (no public link supplied)", document)

    def test_updated_docx_preserves_body_and_footer_images(self):
        import fitz
        from docx import Document
        from docx.shared import Inches
        from xml.etree import ElementTree

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            image = root / "mark.png"
            image_doc = fitz.open()
            image_page = image_doc.new_page(width=120, height=48)
            image_page.draw_rect(image_page.rect, color=(0, 0, 0), fill=(0.85, 0.85, 0.85))
            image_page.insert_text((12, 30), "EXHIBIT", fontsize=16)
            image_page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(image)
            image_doc.close()

            source = root / "brief.docx"
            document = Document()
            document.add_paragraph("R v Oakes, [1986] 1 SCR 103.")
            document.add_picture(str(image), width=Inches(1.5))
            document.sections[0].footer.paragraphs[0].add_run().add_picture(
                str(image),
                width=Inches(0.8),
            )
            document.save(source)
            build_project(
                source,
                root / "out",
                offline=True,
                pdf_mode="none",
                output_mode="table",
            )
            updated = root / "out" / "brief.annotated.docx"
            with zipfile.ZipFile(source) as before, zipfile.ZipFile(updated) as after:
                media = {
                    name: before.read(name)
                    for name in before.namelist()
                    if name.startswith("word/media/")
                }
                self.assertEqual(
                    media,
                    {name: after.read(name) for name in media},
                )
                for part in (
                    "word/_rels/document.xml.rels",
                    "word/_rels/footer1.xml.rels",
                ):
                    self.assertEqual(before.read(part), after.read(part))
                source_root = ElementTree.fromstring(before.read("word/document.xml"))
                updated_root = ElementTree.fromstring(after.read("word/document.xml"))
                count = lambda root: sum(
                    node.tag.rsplit("}", 1)[-1] == "blip"
                    for node in root.iter()
                )
                self.assertEqual(count(source_root), count(updated_root))
                self.assertIn(b" TA \\l ", after.read("word/document.xml"))

    def test_generated_table_pdf_is_monochrome_except_source_link(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "table.pdf"
            authority = Authority(
                "key",
                "case",
                "2020 SCC 1",
                "Example v Example",
                tab="Tab 1",
                source_url="https://example.test/case",
            )
            write_table_of_authorities_pdf([authority], output, "brief.docx")
            with fitz.open(output) as document:
                for page in document:
                    for drawing in page.get_drawings():
                        for color in (drawing.get("color"), drawing.get("fill")):
                            if color is not None:
                                self.assertAlmostEqual(color[0], color[1], places=4)
                                self.assertAlmostEqual(color[1], color[2], places=4)
                    for block in page.get_text("dict")["blocks"]:
                        for line in block.get("lines", []):
                            for span in line["spans"]:
                                red, green, blue = fitz.sRGB_to_rgb(span["color"])
                                if span["text"] == "OPEN":
                                    continue
                                self.assertEqual(red, green)
                                self.assertEqual(green, blue)

    def test_manual_target_is_resolved(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            _docx_with_footnote(source, "R v Oakes, [1986] 1 SCR 103; Ibid.")
            review = review_document(source)
            review.parts[1].supra_target = review.parts[0].part_id
            result = build_project(source, root / "out", review=review, offline=True, pdf_mode="none")
            self.assertEqual(len(result.authorities), 1)
            self.assertEqual(len(result.authorities[0].occurrences), 2)

    def test_different_pinpoints_share_one_authority(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            _docx_with_footnote(source, "R v X, 2020 SCC 1 at para 20; R v X, 2020 SCC 1 at para 30.")
            result = build_project(source, root / "out", offline=True, pdf_mode="none")
            self.assertEqual(len(result.authorities), 1)
            self.assertEqual(
                [occurrence.pinpoint_fragments for occurrence in result.authorities[0].occurrences],
                [["par20"], ["par30"]],
            )

    def test_ibid_is_an_occurrence_of_its_parent(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            _docx_with_footnote(source, "R v Oakes, [1986] 1 SCR 103 at 136; Ibid at 140.")
            review = review_document(source)
            self.assertEqual(review.parts[1].supra_target, review.parts[0].part_id)
            result = build_project(source, root / "out", review=review, offline=True, pdf_mode="none")
            self.assertEqual(len(result.authorities), 1)
            self.assertEqual([item.page_pinpoints for item in result.authorities[0].occurrences], [[136], [140]])

    def test_pdf_mode_writes_structured_outputs(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            _docx_with_footnote(source, "R v Oakes, [1986] 1 SCR 103.")
            result = build_project(
                source,
                root / "out",
                offline=True,
                pdf_mode="render",
                output_mode="both",
            )
            self.assertEqual(len(result.authorities), 1)
            pdfs = list((root / "out" / "authorities").glob("*.pdf"))
            self.assertEqual(len(pdfs), 1)
            self.assertTrue(pdfs[0].read_bytes().startswith(b"%PDF-"))
            book_pdf = root / "out" / "brief.book-of-authorities.pdf"
            self.assertTrue(book_pdf.is_file())
            import fitz

            with fitz.open(book_pdf) as document:
                labels = [row[1] for row in document.get_toc()]
                self.assertIn("Table of Contents", labels)
                self.assertTrue(any(label.startswith("Tab 1") for label in labels))
                self.assertTrue(document[1].get_links())
                self.assertEqual(document.metadata["title"], "Book of Authorities")
                cover_text = document[0].get_text()
                self.assertIn("Book of Authorities", cover_text)
                self.assertNotIn("brief.docx", cover_text)
                self.assertNotIn("LEGAL REFERENCE BOOK", cover_text)
                self.assertNotIn("structured reconstructions", cover_text)
                self.assertNotIn("PDF placeholders", cover_text)
                for page in document:
                    for drawing in page.get_drawings():
                        for color in (drawing.get("color"), drawing.get("fill")):
                            if color is not None:
                                self.assertAlmostEqual(color[0], color[1], places=4)
                                self.assertAlmostEqual(color[1], color[2], places=4)
                    for block in page.get_text("dict")["blocks"]:
                        for line in block.get("lines", []):
                            for span in line["spans"]:
                                red, green, blue = fitz.sRGB_to_rgb(span["color"])
                                self.assertEqual(red, green)
                                self.assertEqual(green, blue)
            self.assertFalse((root / "out" / "brief.table-of-authorities.pdf").exists())
            self.assertFalse((root / "out" / "brief.book-of-authorities.docx").exists())

    def test_placeholder_slot_accepts_manual_pdf_before_finalization(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            output = root / "out"
            _docx_with_footnote(source, "R v Example, 2099 SCC 999.")
            build_project(source, output, offline=True, pdf_mode="render")
            manifest_path = output / "brief.toa-manifest.json"
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            row = manifest["authorities"][0]
            self.assertEqual(row["pdf_origin"], "placeholder")
            with fitz.open(row["pdf_path"]) as placeholder:
                self.assertIn("SOURCE PDF UNAVAILABLE", placeholder[0].get_text())

            manual_path = root / "manual.pdf"
            manual = fitz.open()
            manual.new_page().insert_text((72, 72), "MANUAL AUTHORITY PDF")
            manual.save(manual_path)
            manual.close()
            staged = attach_manual_pdf(manifest_path, row["key"], manual_path)
            self.assertTrue(staged.is_file())
            updated = json.loads(manifest_path.read_text(encoding="utf-8"))
            self.assertEqual(updated["authorities"][0]["pdf_origin"], "manual")
            combined, placeholders = finalize_manifest_book(manifest_path)
            self.assertEqual(placeholders, 0)
            with fitz.open(combined) as document:
                self.assertIn("MANUAL AUTHORITY PDF", document[-1].get_text())

    def test_finalization_can_omit_unfilled_placeholder_pages(self):
        import fitz

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            output = root / "out"
            _docx_with_footnote(source, "R v Example, 2099 SCC 999.")
            build_project(source, output, offline=True, pdf_mode="render")
            manifest_path = output / "brief.toa-manifest.json"

            combined, placeholders = finalize_manifest_book(manifest_path, omit_placeholders=True)

            self.assertEqual(placeholders, 1)
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            self.assertEqual(manifest["omitted_placeholder_count"], 1)
            with fitz.open(combined) as document:
                text = "\n".join(page.get_text() for page in document)
                self.assertNotIn("PDF REQUIRED", text)
                self.assertNotIn("R v Example", text)

    def test_optional_front_matter_and_unrelated_pdf_are_used_on_finalization(self):
        import fitz

        def make_pdf(path, text):
            document = fitz.open()
            document.new_page().insert_text((72, 72), text)
            document.save(path)
            document.close()

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            output = root / "out"
            _docx_with_footnote(source, "R v Example, 2099 SCC 999.")
            build_project(source, output, offline=True, pdf_mode="render")
            manifest_path = output / "brief.toa-manifest.json"
            cover = root / "cover.pdf"
            index = root / "index.pdf"
            extra = root / "extra.pdf"
            make_pdf(cover, "CUSTOM COVER")
            make_pdf(index, "CUSTOM INDEX")
            make_pdf(extra, "UNRELATED EXHIBIT")
            attach_manifest_extra_pdf(manifest_path, "cover", cover)
            attach_manifest_extra_pdf(manifest_path, "index", index)
            attach_manifest_extra_pdf(
                manifest_path,
                "supplemental",
                extra,
                title="Independent exhibit",
                tab="Exhibit X",
            )

            combined, _placeholders = finalize_manifest_book(manifest_path)

            with fitz.open(combined) as document:
                self.assertIn("CUSTOM COVER", document[0].get_text())
                self.assertIn("CUSTOM INDEX", document[1].get_text())
                self.assertIn("UNRELATED EXHIBIT", document[-1].get_text())
                labels = [row[1] for row in document.get_toc()]
                self.assertIn("Exhibit X — Independent exhibit", labels)

    def test_scc_landing_page_uses_native_document_pdf(self):
        class Response:
            status_code = 200
            headers = {"content-type": "application/pdf"}
            encoding = "utf-8"

            def __init__(self, data):
                self.data = data

            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def iter_content(self, chunk_size):
                yield self.data

        class Session:
            def __init__(self, data):
                self.data = data
                self.urls = []

            def get(self, url, **_):
                self.urls.append(url)
                return Response(self.data)

        import fitz

        source = fitz.open()
        source.new_page()
        pdf_data = source.tobytes()
        source.close()
        session = Session(pdf_data)
        with tempfile.TemporaryDirectory() as temp:
            authority = Authority(
                "scc",
                "case",
                "2023 SCC 14",
                "Hansman v. Neufeld",
                source_url="https://decisions.scc-csc.ca/scc-csc/scc-csc/en/item/19911/index.do",
                tab="Tab 1",
            )
            _download_pdf(authority, Path(temp), "auto", session)
            self.assertEqual(
                session.urls[0],
                "https://decisions.scc-csc.ca/scc-csc/scc-csc/en/19911/1/document.do",
            )
            self.assertEqual(authority.pdf_origin, "original")
            self.assertEqual(authority.pdf_source_url, session.urls[0])
            self.assertTrue(Path(authority.pdf_path).is_file())

    def test_canlii_pdf_is_never_requested_automatically(self):
        class Session:
            def __init__(self):
                self.urls = []

            def get(self, url, **_):
                self.urls.append(url)
                raise AssertionError("CanLII must remain a user-initiated download")

        session = Session()
        with tempfile.TemporaryDirectory() as temp:
            authority = Authority(
                "scc",
                "case",
                "2009 SCC 32",
                "R v Grant",
                source_url="https://www.canlii.org/en/ca/scc/doc/2009/2009scc32/2009scc32.html",
                tab="Tab 1",
            )
            _download_pdf(authority, Path(temp), "originals", session)
            self.assertEqual(session.urls, [])
            self.assertEqual(authority.pdf_origin, "placeholder")
            self.assertTrue(Path(authority.pdf_path).is_file())

    def test_tab_marker_follows_authority_not_pinpoint(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "brief.docx"
            _docx_with_footnote(source, "R v Oakes, [1986] 1 SCR 103 at 136.")
            build_project(
                source,
                root / "out",
                offline=True,
                pdf_mode="none",
                output_mode="both",
            )
            with zipfile.ZipFile(root / "out" / "brief.annotated.docx") as archive:
                footnotes = archive.read("word/footnotes.xml").decode("utf-8")
            from xml.etree import ElementTree

            root_xml = ElementTree.fromstring(footnotes)
            visible = "".join(node.text or "" for node in root_xml.iter(f"{{{W}}}t"))
            self.assertIn("103 [Tab 1] at 136", visible)


class CitationSpanTests(unittest.TestCase):
    def test_tight_crop_can_expand_and_pinpoint_highlight_stays_exact(self):
        text = "See R v Oakes, [1986] 1 SCR 103 at 136, discussed later."
        unit = TextUnit("footnote:1", "footnote", 0, 1, text)
        citation_start = text.index("[1986]")
        part_end = text.index(" at 136")
        part = _part_from(
            unit,
            1,
            "manual",
            (),
            DeterministicPart(citation_start, part_end, text[citation_start:part_end], ("reporter",)),
        )
        expanded_start = text.index("R v Oakes")
        _set_authority_span(part, unit, expanded_start, text.index(" at 136"))
        pinpoint_start = text.index("at 136")
        _set_pinpoint_span(part, unit, pinpoint_start, pinpoint_start + len("at 136"))
        self.assertEqual(part.start, expanded_start)
        self.assertEqual(part.authority_text, "R v Oakes, [1986] 1 SCR 103")
        self.assertEqual(part.page_pinpoints, [136])
        self.assertEqual(part.end, pinpoint_start + len("at 136"))


if __name__ == "__main__":
    unittest.main()
